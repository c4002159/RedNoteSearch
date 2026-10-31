"""Parsing, chunking, embedding client and review input rules. None of these need a database."""

import csv
import io
import json

import httpx
import pytest
from docx import Document
from pydantic import ValidationError

from app.fakes.embedding import FakeEmbeddingProvider
from app.infra.dashscope import BATCH_SIZE, DashScopeEmbeddingProvider
from app.models import EMBEDDING_DIMENSIONS, ReviewCreate
from app.ports.embedding import EmbeddingProvider
from app.services.chunking import chunk_text
from app.services.errors import ModelUnavailableError, ParseError
from app.services.parsing import parse_document


def _pdf_with_text(text: str) -> bytes:
    """A minimal one-page PDF whose page shows ``text`` (ASCII only)."""
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(objects) + 1,
        xref,
    )
    return out


def test_chunking_keeps_paragraphs_and_respects_the_limit() -> None:
    assert chunk_text("") == [] and chunk_text("  \n\n ") == []
    assert chunk_text("第一段。\n\n第二段。") == ["第一段。\n第二段。"]
    text = "\n\n".join(f"段落{i}" + "字" * 40 for i in range(20))
    chunks = chunk_text(text, max_chars=100)
    assert len(chunks) > 1 and all(len(c) <= 100 for c in chunks)
    assert "".join(chunks).replace("\n", "") == text.replace("\n\n", "")


def test_chunking_splits_long_paragraphs_at_sentence_ends() -> None:
    sentence = "这是一个完整的句子。"
    paragraph = sentence * 30
    chunks = chunk_text(paragraph, max_chars=50)
    assert all(len(c) <= 50 and c.endswith("。") for c in chunks)
    assert "".join(chunks) == paragraph
    # A sentence longer than the limit is cut hard rather than dropped.
    long_sentence = "无标点" * 40
    hard = chunk_text(long_sentence, max_chars=30)
    assert all(len(c) <= 30 for c in hard) and "".join(hard) == long_sentence


def test_parse_text_markdown_and_encodings() -> None:
    assert parse_document("a.md", "# 标题\n正文".encode()) == "# 标题\n正文"
    assert parse_document("a.TXT", "品牌故事".encode("gb18030")) == "品牌故事"
    assert parse_document("a.txt", b"\xef\xbb\xbfbom text") == "bom text"


def test_parse_csv_turns_rows_into_labelled_lines() -> None:
    buffer = io.StringIO()
    csv.writer(buffer).writerows([["成分", "功效"], ["玫瑰", "香气"], ["", ""]])
    text = parse_document("data.csv", buffer.getvalue().encode())
    assert text == "成分：玫瑰；功效：香气"


def test_parse_docx_reads_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("冰希黎品牌故事")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "前调"
    table.rows[0].cells[1].text = "佛手柑"
    buffer = io.BytesIO()
    document.save(buffer)
    text = parse_document("story.docx", buffer.getvalue())
    assert "冰希黎品牌故事" in text and "前调；佛手柑" in text


def test_parse_pdf_extracts_text() -> None:
    assert "Bergamot" in parse_document("spec.pdf", _pdf_with_text("Bergamot top note"))


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("a.exe", b"x"),
        ("a", b"x"),
        ("a.txt", b""),
        ("a.txt", b"   \n "),
        ("a.txt", b"\xff\xfe\x00\x81"),
        ("a.csv", b""),
        ("a.docx", b"not a zip"),
        ("a.pdf", b"not a pdf"),
        ("a.pdf", _pdf_with_text("")),  # a page without text, like a scanned PDF
    ],
)
def test_parse_failures_raise_parse_error(filename: str, content: bytes) -> None:
    with pytest.raises(ParseError):
        parse_document(filename, content)


def test_fake_embedding_is_deterministic_and_similar_for_similar_text() -> None:
    fake = FakeEmbeddingProvider()
    assert isinstance(fake, EmbeddingProvider)
    a, b, c = fake.embed(["佛手柑前调", "佛手柑前调", "完全无关的话题"])
    assert a == b and len(a) == EMBEDDING_DIMENSIONS

    def dot(x: list[float], y: list[float]) -> float:
        return sum(i * j for i, j in zip(x, y, strict=True))

    assert dot(a, a) > dot(a, c)
    with pytest.raises(ModelUnavailableError):
        FakeEmbeddingProvider(fail=True).embed(["x"])


def _provider(
    handler: httpx.MockTransport, key: str = "test-key"
) -> DashScopeEmbeddingProvider:
    return DashScopeEmbeddingProvider(
        api_key=key,
        base_url="https://example.test/v1/",
        model="text-embedding-v4",
        client=httpx.Client(transport=handler),
    )


def test_dashscope_batches_and_sends_explicit_dimensions() -> None:
    seen: list[dict[str, object]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        seen.append(body)
        assert request.url.path == "/v1/embeddings"
        assert request.headers["authorization"] == "Bearer test-key"
        data = [  # answered out of order on purpose
            {"index": i, "embedding": [float(i)] + [0.0] * (EMBEDDING_DIMENSIONS - 1)}
            for i in reversed(range(len(body["input"])))
        ]
        return httpx.Response(200, json={"data": data})

    texts = [f"t{i}" for i in range(BATCH_SIZE + 3)]
    vectors = _provider(httpx.MockTransport(handler)).embed(texts)
    assert [len(s["input"]) for s in seen] == [BATCH_SIZE, 3]  # type: ignore[arg-type]
    assert all(s["dimensions"] == EMBEDDING_DIMENSIONS for s in seen)
    assert [v[0] for v in vectors] == [float(i) for i in range(BATCH_SIZE)] + [
        0.0,
        1.0,
        2.0,
    ]


def test_dashscope_reports_failures_without_leaking_details() -> None:
    def server_error(_: httpx.Request) -> httpx.Response:
        return httpx.Response(500, text="secret-key echoed here")

    def wrong_size(_: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"data": [{"index": 0, "embedding": [1.0]}]})

    def garbage(_: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"oops": 1})

    for handler in (server_error, wrong_size, garbage):
        with pytest.raises(ModelUnavailableError) as error:
            _provider(httpx.MockTransport(handler)).embed(["x"])
        assert "secret" not in error.value.message
    with pytest.raises(ModelUnavailableError, match="not configured"):
        _provider(httpx.MockTransport(server_error), key="").embed(["x"])


def test_dashscope_network_errors_become_model_unavailable() -> None:
    def refuse(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("down", request=request)

    with pytest.raises(ModelUnavailableError):
        _provider(httpx.MockTransport(refuse)).embed(["x"])


def test_reject_review_requires_a_reason() -> None:
    assert ReviewCreate(action="approve").reason is None
    assert ReviewCreate(action="reject", reason="来源不明").action == "reject"
    for bad in (
        {"action": "reject"},
        {"action": "reject", "reason": "  "},
        {"action": "delete"},
    ):
        with pytest.raises(ValidationError):
            ReviewCreate.model_validate(bad)
